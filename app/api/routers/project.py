from datetime import datetime, timezone, date
from typing import Optional, List
import math
import io
import urllib.parse
import os

import httpx
from dotenv import load_dotenv
from sqlalchemy.orm import Session

from fastapi import APIRouter, Depends, HTTPException, Query, status, BackgroundTasks
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

from hwpx.document import HwpxDocument
from bs4 import BeautifulSoup

from app.db.session import get_db
from app.models.project import Project
from app.models.project_member import ProjectMember
from app.models.user import User
from app.models.department import Department
from app.api.routers.auth import get_current_user

FONT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "..", "fonts")
FONT_DIR = os.path.normpath(FONT_DIR)
pdfmetrics.registerFont(TTFont('MalgunGothic', os.path.join(FONT_DIR, 'malgun.ttf')))
pdfmetrics.registerFont(TTFont('MalgunGothicBold', os.path.join(FONT_DIR, 'malgunbd.ttf')))
pdfmetrics.registerFontFamily(
    'MalgunGothic',
    normal='MalgunGothic',
    bold='MalgunGothicBold',
    italic='MalgunGothic',
    boldItalic='MalgunGothicBold',
)

load_dotenv(".env.local")
AI_SERVER = os.getenv("AI_SERVER")

router = APIRouter()

# ----------------------------------------------------------------
# 상수
# ----------------------------------------------------------------

STAGE_NAME_MAP = {
    "01": "저장",
    "02": "승인완료",
}

ROLE_NAME_MAP = {
    "01": "주관",
    "02": "협력",
}

# ----------------------------------------------------------------
# Pydantic 스키마
# ----------------------------------------------------------------

class ProjectCreateRequest(BaseModel):
    name: str
    businessContent: str
    startDate: Optional[str] = None
    deadline: Optional[str] = None
    memberUserIds: List[int]

class ProjectCreateResponse(BaseModel):
    projectId: int

class ProjectUpdateRequest(BaseModel):
    name: str
    businessContent: str
    startDate: Optional[str] = None
    deadline: Optional[str] = None
    overview: Optional[str] = None
    reportContent: Optional[str] = None
    secOverview: Optional[str] = None
    secBackground: Optional[str] = None
    secGoals: Optional[str] = None
    secDetailedPlan: Optional[str] = None
    secSchedule: Optional[str] = None
    secExecutionSystem: Optional[str] = None
    secBudget: Optional[str] = None
    secExpectedEffect: Optional[str] = None
    secPostManagement: Optional[str] = None
    coverTitle: Optional[str] = None
    memberUserIds: List[int]

class MemberItem(BaseModel):
    userId: int
    name: str
    roleName: str
    departmentName: str

class ProjectListItem(BaseModel):
    projectId: int
    name: str
    stageName: str
    departmentName: str
    startDate: Optional[str]
    deadline: Optional[str]
    createdAt: str
    roleType: str

class ProjectListResponse(BaseModel):
    content: List[ProjectListItem]
    totalElements: int
    totalPages: int
    page: int
    size: int

class ProjectDetailResponse(BaseModel):
    projectId: int
    name: str
    stageCode: str
    stageName: str
    departmentName: str
    startDate: Optional[str]
    deadline: Optional[str]
    businessContent: Optional[str]
    overview: Optional[str]
    reportContent: Optional[str]
    secOverview: Optional[str]
    secBackground: Optional[str]
    secGoals: Optional[str]
    secDetailedPlan: Optional[str]
    secSchedule: Optional[str]
    secExecutionSystem: Optional[str]
    secBudget: Optional[str]
    secExpectedEffect: Optional[str]
    secPostManagement: Optional[str]
    coverTitle: Optional[str]
    approvedAt: Optional[str]
    createdAt: str
    members: List[MemberItem]

class AiDraftResponse(BaseModel):
    draft: dict
    referenced_tasks: list = []
    guardrail_triggered: bool = False
    needs_review: bool = False
    unverified_claims: dict = {}

# ----------------------------------------------------------------
# 유틸
# ----------------------------------------------------------------

def date_to_str(d) -> Optional[str]:
    return d.isoformat() if d else None

def datetime_to_str(dt) -> Optional[str]:
    return dt.isoformat() if dt else None

def str_to_date(s: Optional[str]) -> Optional[date]:
    if not s:
        return None
    return date.fromisoformat(s)

def get_department_name(department_code: Optional[str], db: Session) -> str:
    if not department_code:
        return ""
    dept = db.query(Department).filter(
        Department.department_code == department_code
    ).first()
    return dept.name if dept else ""

def build_detail_response(project: Project, db: Session) -> ProjectDetailResponse:
    members_raw = db.query(ProjectMember, User).join(
        User, ProjectMember.user_id == User.user_id
    ).filter(ProjectMember.project_id == project.project_id).all()

    members = [
        MemberItem(
            userId=u.user_id,
            name=u.name,
            roleName=ROLE_NAME_MAP.get(pm.role_code, ""),
            departmentName=get_department_name(u.department_code, db),
        )
        for pm, u in members_raw
    ]

    return ProjectDetailResponse(
        projectId=project.project_id,
        name=project.name,
        stageCode=project.stage_code or "",
        stageName=STAGE_NAME_MAP.get(project.stage_code, ""),
        departmentName=get_department_name(project.department_code, db),
        startDate=date_to_str(project.start_date),
        deadline=date_to_str(project.deadline),
        businessContent=project.business_content,
        overview=project.overview,
        reportContent=project.report_content,
        secOverview=project.sec_overview,
        secBackground=project.sec_background,
        secGoals=project.sec_goals,
        secDetailedPlan=project.sec_detailed_plan,
        secSchedule=project.sec_schedule,
        secExecutionSystem=project.sec_execution_system,
        secBudget=project.sec_budget,
        secExpectedEffect=project.sec_expected_effect,
        secPostManagement=project.sec_post_management,
        coverTitle=project.cover_title,
        approvedAt=datetime_to_str(project.approved_at),
        createdAt=datetime_to_str(project.created_at),
        members=members,
    )

# ----------------------------------------------------------------
# 엔드포인트
# ----------------------------------------------------------------

# 1. 프로젝트 목록 조회
@router.get(
    "",
    response_model=ProjectListResponse,
    summary="프로젝트 목록 조회",
)
def get_project_list(
    scope: str = Query(..., description="MY(내 주관) | JOINED(내 참여) | PREDECESSOR(전임자)"),
    stage: Optional[str] = Query(None),
    page: int = Query(0, ge=0),
    size: int = Query(10, ge=1),
    keyword: Optional[str] = Query(None),
    department_code: Optional[str] = Query(None),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    # 관리자는 전체 프로젝트 조회
    if current_user.system_role_code == "02":
        query = db.query(Project)

    elif scope == "MY":
        project_ids = db.query(ProjectMember.project_id).filter(
            ProjectMember.user_id == current_user.user_id,
            ProjectMember.role_code == "01",
        ).subquery()
        query = db.query(Project).filter(Project.project_id.in_(project_ids))

    elif scope == "JOINED":
        project_ids = db.query(ProjectMember.project_id).filter(
            ProjectMember.user_id == current_user.user_id,
            ProjectMember.role_code == "02",
        ).subquery()
        query = db.query(Project).filter(Project.project_id.in_(project_ids))

    elif scope == "PREDECESSOR":
        if not current_user.predecessor_user_id:
            return ProjectListResponse(content=[], totalElements=0, totalPages=0, page=page, size=size)
        project_ids = db.query(ProjectMember.project_id).filter(
            ProjectMember.user_id == current_user.predecessor_user_id,
            ProjectMember.role_code == "01",
        ).subquery()
        query = db.query(Project).filter(
            Project.project_id.in_(project_ids),
            Project.stage_code == "02",
        )

    else:
        raise HTTPException(status_code=400, detail="유효하지 않은 scope 값입니다.")

    if stage:
        query = query.filter(Project.stage_code == stage)
    
    if keyword:
        query = query.filter(Project.name.like(f"%{keyword}%"))

    if department_code:
        query = query.filter(Project.department_code == department_code)

    total_elements = query.count()
    total_pages = math.ceil(total_elements / size)
    items = query.order_by(Project.start_date.asc()).offset(page * size).limit(size).all()

    content = [
        ProjectListItem(
            projectId=p.project_id,
            name=p.name,
            stageName=STAGE_NAME_MAP.get(p.stage_code, ""),
            departmentName=get_department_name(p.department_code, db),
            startDate=date_to_str(p.start_date),
            deadline=date_to_str(p.deadline),
            createdAt=datetime_to_str(p.created_at),
            roleType="ADMIN" if current_user.system_role_code == "02" else scope,
        )
        for p in items
    ]

    return ProjectListResponse(
        content=content,
        totalElements=total_elements,
        totalPages=total_pages,
        page=page,
        size=size,
    )


# 2. 새 프로젝트 생성
@router.post(
    "",
    response_model=ProjectCreateResponse,
    status_code=status.HTTP_200_OK,
    summary="새 프로젝트 생성",
)
def create_project(
    body: ProjectCreateRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    project = Project(
        name=body.name,
        business_content=body.businessContent,
        start_date=str_to_date(body.startDate),
        deadline=str_to_date(body.deadline),
        department_code=current_user.department_code,
        stage_code="01",
    )
    db.add(project)
    db.flush()

    owner_member = ProjectMember(
        project_id=project.project_id,
        user_id=current_user.user_id,
        role_code="01",
        invited_by=current_user.user_id,
    )
    db.add(owner_member)

    for uid in body.memberUserIds:
        if uid == current_user.user_id:
            continue
        member = ProjectMember(
            project_id=project.project_id,
            user_id=uid,
            role_code="02",
            invited_by=current_user.user_id,
        )
        db.add(member)

    db.commit()
    db.refresh(project)

    return ProjectCreateResponse(projectId=project.project_id)


# 3. 프로젝트 상세 조회
@router.get(
    "/{projectId}",
    response_model=ProjectDetailResponse,
    summary="프로젝트 상세 조회",
)
def get_project_detail(
    projectId: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    project = db.query(Project).filter(Project.project_id == projectId).first()
    if not project:
        raise HTTPException(status_code=404, detail="존재하지 않는 프로젝트입니다.")

    return build_detail_response(project, db)


# 4. 프로젝트 저장
@router.patch(
    "/{projectId}",
    response_model=ProjectDetailResponse,
    summary="프로젝트 저장",
)
def update_project(
    projectId: int,
    body: ProjectUpdateRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    project = db.query(Project).filter(Project.project_id == projectId).first()
    if not project:
        raise HTTPException(status_code=404, detail="존재하지 않는 프로젝트입니다.")

    if project.stage_code == "02":
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="승인완료 상태라 수정할 수 없습니다.",
        )

    project.name = body.name
    project.business_content = body.businessContent
    project.start_date = str_to_date(body.startDate)
    project.deadline = str_to_date(body.deadline)
    if body.overview is not None:
        project.overview = body.overview
    if body.reportContent is not None:
        project.report_content = body.reportContent
    if body.secOverview is not None:
        project.sec_overview = body.secOverview
    if body.secBackground is not None:
        project.sec_background = body.secBackground
    if body.secGoals is not None:
        project.sec_goals = body.secGoals
    if body.secDetailedPlan is not None:
        project.sec_detailed_plan = body.secDetailedPlan
    if body.secSchedule is not None:
        project.sec_schedule = body.secSchedule
    if body.secExecutionSystem is not None:
        project.sec_execution_system = body.secExecutionSystem
    if body.secBudget is not None:
        project.sec_budget = body.secBudget
    if body.secExpectedEffect is not None:
        project.sec_expected_effect = body.secExpectedEffect
    if body.secPostManagement is not None:
        project.sec_post_management = body.secPostManagement
    if body.coverTitle is not None:
        project.cover_title = body.coverTitle

    existing_members = db.query(ProjectMember).filter(
        ProjectMember.project_id == projectId
    ).all()

    existing_collab_ids = {
        m.user_id for m in existing_members if m.role_code == "02"
    }
    new_collab_ids = set(body.memberUserIds)

    ids_to_remove = existing_collab_ids - new_collab_ids
    if ids_to_remove:
        db.query(ProjectMember).filter(
            ProjectMember.project_id == projectId,
            ProjectMember.user_id.in_(ids_to_remove),
            ProjectMember.role_code == "02",
        ).delete(synchronize_session=False)

    ids_to_add = new_collab_ids - existing_collab_ids
    for uid in ids_to_add:
        db.add(ProjectMember(
            project_id=projectId,
            user_id=uid,
            role_code="02",
            invited_by=current_user.user_id,
        ))

    db.commit()
    db.refresh(project)

    return build_detail_response(project, db)

# AI 서버 색인 등록 (백그라운드 실행용) - 승인완료 시 호출
def _call_index_task(project: Project):
    try:
        httpx.post(
            f"{AI_SERVER}/api/index-task",
            json={
                "task_id": project.project_id,
                "year": project.approved_at.year,
                "title": project.name,
                "lead_department_code": project.department_code or "01",
                "collab_department_codes": [],
                "domain_code": "",
                "overview": project.sec_overview or "",
                "background": project.sec_background or "",
                "goals": project.sec_goals or "",
                "detailed_plan": project.sec_detailed_plan or "",
                "schedule": project.sec_schedule or "",
                "execution_system": project.sec_execution_system or "",
                "budget": project.sec_budget or "",
                "expected_effect": project.sec_expected_effect or "",
                "post_management": project.sec_post_management or "",
                "status_code": "완료",
            },
            timeout=180.0,
        )
    except Exception:
        pass


# 5. 기획서 승인완료
@router.post(
    "/{projectId}/approve",
    status_code=status.HTTP_200_OK,
    summary="기획서 승인완료",
)
def approve_project(
    projectId: int,
    background_tasks: BackgroundTasks,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    project = db.query(Project).filter(Project.project_id == projectId).first()
    if not project:
        raise HTTPException(status_code=404, detail="존재하지 않는 프로젝트입니다.")

    if project.stage_code == "02":
        raise HTTPException(status_code=409, detail="이미 승인완료 상태입니다.")

    project.stage_code = "02"
    project.approved_at = datetime.now(timezone.utc)
    db.commit()

    background_tasks.add_task(_call_index_task, project)

    return {"message": "승인완료 처리되었습니다."}

# 6. 프로젝트 삭제
@router.delete(
    "/{projectId}",
    status_code=status.HTTP_200_OK,
    summary="프로젝트 삭제",
)
def delete_project(
    projectId: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    project = db.query(Project).filter(Project.project_id == projectId).first()
    if not project:
        raise HTTPException(status_code=404, detail="존재하지 않는 프로젝트입니다.")

    if project.stage_code != "01":
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="저장(기획중) 상태인 프로젝트만 삭제할 수 있습니다.",
        )

    db.query(ProjectMember).filter(
        ProjectMember.project_id == projectId
    ).delete(synchronize_session=False)

    db.delete(project)
    db.commit()


# 7. AI 기획서 초안 생성
@router.post(
    "/{projectId}/ai-draft",
    response_model=AiDraftResponse,
    summary="AI 기획서 초안 생성",
)
def get_ai_draft(
    projectId: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    project = db.query(Project).filter(Project.project_id == projectId).first()
    if not project:
        raise HTTPException(status_code=404, detail="존재하지 않는 프로젝트입니다.")

    VALID_DEPT_CODES = {"01","02","03","04","05","06","07","08"}
    dept_code = project.department_code if project.department_code in VALID_DEPT_CODES else "01"

    try:
        payload = {
            "title": project.name,
            "overview": project.overview or project.business_content or "",
            "lead_department_code": dept_code,
        }
        res = httpx.post(
            f"{AI_SERVER}/api/task-draft",
            json=payload,
            timeout=280.0,
        )
        res.raise_for_status()
        data = res.json()
    except httpx.TimeoutException:
        raise HTTPException(status_code=504, detail="AI 서버 응답 시간이 초과되었습니다. 잠시 후 다시 시도해주세요.")
    except httpx.HTTPStatusError as e:
        raise HTTPException(status_code=502, detail=f"AI 서버 오류: {e.response.status_code}")
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"AI 서버 연결 오류: {str(e)}")

    return AiDraftResponse(
        draft=data.get("draft", {}),
        referenced_tasks=data.get("referenced_tasks", []),
        guardrail_triggered=data.get("guardrail_triggered", False),
        needs_review=data.get("needs_review", False),
        unverified_claims=data.get("unverified_claims", {}),
    )

# 8. 기획서 내보내기
@router.get(
    "/{projectId}/export",
    summary="기획서 내보내기",
)
def export_project(
    projectId: int,
    format: str = Query(..., description="pdf | docx | hwpx"),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    project = db.query(Project).filter(Project.project_id == projectId).first()
    if not project:
        raise HTTPException(status_code=404, detail="존재하지 않는 프로젝트입니다.")

    title = project.name or "기획서"

    # 9개 섹션 정의
    SECTIONS = [
        ("Ⅰ. 사업 개요",           project.sec_overview),
        ("Ⅱ. 추진 배경 및 필요성",  project.sec_background),
        ("Ⅲ. 사업 목표",           project.sec_goals),
        ("Ⅳ. 세부 추진 계획",       project.sec_detailed_plan),
        ("Ⅴ. 추진 일정",           project.sec_schedule),
        ("Ⅵ. 사업 추진 체계",       project.sec_execution_system),
        ("Ⅶ. 예산 계획",           project.sec_budget),
        ("Ⅷ. 기대 효과",           project.sec_expected_effect),
        ("Ⅸ. 사후 관리 계획",       project.sec_post_management),
    ]

    # ── 표지 데이터 준비 ──────────────────────────────────
    members_raw = db.query(ProjectMember, User).join(
        User, ProjectMember.user_id == User.user_id
    ).filter(ProjectMember.project_id == projectId).all()

    owner = next((u for pm, u in members_raw if pm.role_code == "01"), None)
    owner_name = owner.name if owner else ""
    dept_name = get_department_name(project.department_code, db)
    year = project.start_date.year if project.start_date else datetime.now().year
    cover_title_text = project.cover_title or f"{project.name} 추진계획서"
    created_date = project.created_at.strftime("%Y. %m. %d.") if project.created_at else ""
    # ────────────────────────────────────────────────────

    def html_to_text(html: str) -> str:
        if not html:
            return ""
        soup = BeautifulSoup(html, "html.parser")
        for tag in soup.find_all("h2"):
            tag.decompose()
        for br in soup.find_all("br"):
            br.replace_with("\n")
        lines = []
        for elem in soup.find_all(["h3", "p"]):
            text = elem.get_text(separator="", strip=False)
            text = text.replace("\xa0", " ")
            text = text.strip()
            if text.strip():
                lines.append(text)
            else:
                lines.append("")
        return "\n".join(lines)
    
    # ── DOCX ──────────────────────────────────────────────
    if format == "docx":
        doc = DocxDocument()
        style = doc.styles['Normal']
        style.font.name = '맑은 고딕'
        style.element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')

        def add_heading(text, level):
            h = doc.add_heading(text, level=level)
            for run in h.runs:
                run.font.name = '맑은 고딕'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')

        def add_paragraph(text):
            p = doc.add_paragraph(text)
            for run in p.runs:
                run.font.name = '맑은 고딕'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')

        # ── 표지 (1페이지) ──
        def add_cover_para(text, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER):
            p = doc.add_paragraph()
            p.alignment = align
            run = p.add_run(text)
            run.font.name = '맑은 고딕'
            run.font.size = Pt(size)
            run.font.bold = bold
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')

        def add_border_para(text, size=12, bold=False, border_top=False, border_bottom=False, space_before=0):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(space_before)
            run = p.add_run(text)
            run.font.name = '맑은 고딕'
            run.font.size = Pt(size)
            run.font.bold = bold
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
            pPr = p._element.get_or_add_pPr()
            pBdr = OxmlElement('w:pBdr')
            if border_top:
                top = OxmlElement('w:top')
                top.set(qn('w:val'), 'single')
                top.set(qn('w:sz'), '24')
                top.set(qn('w:space'), '4')
                top.set(qn('w:color'), '0078D7')
                pBdr.append(top)
            if border_bottom:
                bottom = OxmlElement('w:bottom')
                bottom.set(qn('w:val'), 'single')
                bottom.set(qn('w:sz'), '24')
                bottom.set(qn('w:space'), '4')
                bottom.set(qn('w:color'), '0078D7')
                pBdr.append(bottom)
            pPr.append(pBdr)

        add_cover_para("")
        add_cover_para("")

        # 위 파란선 + 공백
        add_border_para("", size=2, border_top=True, space_before=2)
        # 연도
        add_cover_para(f"{year}년도", size=13, bold=True)
        # 제목
        add_cover_para(cover_title_text, size=20, bold=True)
        # 공백 + 아래 파란선
        add_border_para("", size=2, border_bottom=True, space_before=2)
        for _ in range(7):
            add_cover_para("")

        # 정보 항목
        info_items_docx = [
            ("사  업  명", project.name),
            ("담 당 부 서", dept_name),
            ("작  성  자", owner_name),
            ("작  성  일", created_date),
        ]
        
        for label, value in info_items_docx:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Pt(120)

            # 탭스톱 설정 (콜론 위치 고정)
            pPr = p._element.get_or_add_pPr()
            tabs = OxmlElement('w:tabs')
            tab = OxmlElement('w:tab')
            tab.set(qn('w:val'), 'left')
            tab.set(qn('w:pos'), '1440')
            tabs.append(tab)
            pPr.append(tabs)

            # 라벨
            run_label = p.add_run(label)
            run_label.font.name = '맑은 고딕'
            run_label.font.size = Pt(11)
            run_label._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
            # 탭 → 콜론 → 값
            run_colon = p.add_run("\t:  ")
            run_colon.font.name = '맑은 고딕'
            run_colon.font.size = Pt(11)
            run_colon._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
            # 값
            run_value = p.add_run(value)
            run_value.font.name = '맑은 고딕'
            run_value.font.size = Pt(11)
            run_value._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')

        # 2페이지부터 본문
        doc.add_page_break()

        for sec_title, sec_content in SECTIONS:
            if sec_content:
                add_heading(sec_title, level=2)
                add_paragraph(html_to_text(sec_content))

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f"attachment; filename*=UTF-8''{urllib.parse.quote(title)}.docx"}
        )

    # ── PDF ───────────────────────────────────────────────
    elif format == "pdf":
        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=A4)
        width, height = A4

        def new_page_if_needed(y, needed=30):
            if y < needed:
                c.showPage()
                c.setFont("MalgunGothic", 11)
                return height - 50
            return y

        # ── 표지 (1페이지) ──
        # 제목 줄바꿈 함수
        def wrap_text(text, font, size, max_width):
            lines = []
            current = ""
            for char in text:
                test = current + char
                if c.stringWidth(test, font, size) > max_width:
                    lines.append(current)
                    current = char
                else:
                    current = test
            if current:
                lines.append(current)
            return lines

        # ── 제목 블록
        block_top = height - (height * 0.15)

        # 위 파란 선
        c.setStrokeColorRGB(0.0, 0.47, 0.84)
        c.setLineWidth(5)
        c.line(width * 0.08, block_top, width * 0.92, block_top)

        # 연도
        c.setFillColorRGB(0, 0, 0)
        c.setFont("MalgunGothicBold", 12)
        c.drawCentredString(width / 2, block_top - 64, f"{year}년도")

        # 제목 줄바꿈
        title_lines = wrap_text(cover_title_text, "MalgunGothicBold", 22, width * 0.76)
        title_y = block_top - 106
        c.setFont("MalgunGothicBold", 22)
        for line in title_lines:
            c.drawCentredString(width / 2, title_y, line)
            title_y -= 30

        # 아래 파란 선
        block_bottom = title_y - 30
        c.setStrokeColorRGB(0.0, 0.47, 0.84)
        c.setLineWidth(5)
        c.line(width * 0.08, block_bottom, width * 0.92, block_bottom)

        # ── 정보 항목 (구분선 아래, 중앙 정렬 테이블)
        info_items = [
            ("사  업  명", project.name),
            ("담 당 부 서", dept_name),
            ("작  성  자", owner_name),
            ("작  성  일", created_date),
        ]

        info_y = height * 0.22
        line_gap = 26
        colon_x = width * 0.40
        value_x = width * 0.45

        c.setFont("MalgunGothic", 11)
        for label, value in info_items:
            c.drawRightString(colon_x - 6, info_y, label)  # 라벨 오른쪽 정렬
            c.drawString(colon_x, info_y, ":")
            c.drawString(value_x + 6, info_y, value)
            info_y -= line_gap

        # 2페이지부터 본문
        c.showPage()
        y = height - 50

        for sec_title, sec_content in SECTIONS:
            if not sec_content:
                continue
            # 섹션 제목
            y = new_page_if_needed(y, 60)
            c.setFont("MalgunGothicBold", 13)
            c.drawString(50, y, sec_title)
            y -= 25

            # 섹션 내용
            c.setFont("MalgunGothic", 11)
            max_width = width - 120
            for line in html_to_text(sec_content).split("\n"):
                line = line.strip()
                if not line:
                    y -= 8
                    continue
                # 너비 기반 줄바꿈
                while line:
                    y = new_page_if_needed(y)
                    cut = len(line)
                    while cut > 0 and c.stringWidth(line[:cut], "MalgunGothic", 11) > max_width:
                        cut -= 1
                    if cut == 0:
                        cut = 1
                    c.drawString(60, y, line[:cut])
                    line = line[cut:]
                    y -= 16

            y -= 12  # 섹션 간격

        c.save()
        buf.seek(0)
        return StreamingResponse(
            buf,
            media_type="application/pdf",
            headers={"Content-Disposition": f"attachment; filename*=UTF-8''{urllib.parse.quote(title)}.pdf"}
        )

    # ── HWPX ──────────────────────────────────────────────
    elif format == "hwpx":
        doc = HwpxDocument.new()
        # ── 표지 (1페이지) ──
        doc.add_paragraph("")
        doc.add_paragraph(f"{year}년도")
        doc.add_paragraph("")
        doc.add_paragraph(cover_title_text)
        doc.add_paragraph("")
        doc.add_paragraph(f"사업명    :  {project.name}")
        doc.add_paragraph(f"담당부서  :  {dept_name}")
        doc.add_paragraph(f"작성자    :  {owner_name}")
        doc.add_paragraph(f"작성일    :  {created_date}")
        doc.add_paragraph("\x0C")

        for sec_title, sec_content in SECTIONS:
            if sec_content:
                doc.add_paragraph(sec_title)
                doc.add_paragraph(html_to_text(sec_content))

        buf = io.BytesIO()
        doc.save_to_stream(buf)
        buf.seek(0)
        return StreamingResponse(
            buf,
            media_type="application/hwpx",
            headers={"Content-Disposition": f"attachment; filename*=UTF-8''{urllib.parse.quote(title)}.hwpx"}
        )

    else:
        raise HTTPException(status_code=400, detail="지원하지 않는 형식입니다.")